# Copyright (c) 2023, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
import json
from frappe import _
from frappe.desk.form.assign_to import add as assign
from frappe.utils.file_manager import get_file_path
from frappe.model.mapper import get_mapped_doc
from frappe.utils import now_datetime
from frappe.contacts.doctype.address.address import get_address_display

from crm.fcrm.doctype.crm_service_level_agreement.utils import get_sla
from crm.fcrm.doctype.crm_status_change_log.crm_status_change_log import add_status_change_log


class CRMOrganization(Document):
	@staticmethod
	def default_list_data():
		columns = [
			{
				'label': 'Organization',
				'type': 'Data',
				'key': 'organization_name',
				'width': '16rem',
			},
			{
				'label': 'Website',
				'type': 'Data',
				'key': 'website',
				'width': '14rem',
			},
			{
				'label': 'Industry',
				'type': 'Link',
				'key': 'industry',
				'options': 'CRM Industry',
				'width': '14rem',
			},
			{
				'label': 'Annual Revenue',
				'type': 'Currency',
				'key': 'annual_revenue',
				'width': '14rem',
			},
			{
				'label': 'Last Modified',
				'type': 'Datetime',
				'key': 'modified',
				'width': '8rem',
			},
		]
		rows = [
			"name",
			"organization_name",
			"organization_logo",
			"website",
			"industry",
			"currency",
			"annual_revenue",
			"modified",
		]
		return {'columns': columns, 'rows': rows}


class CRMOrganization(Document):
	def before_insert(self):
		if self.is_new():
			self.set_naming_series()
	def set_naming_series(self):
		if not self.organization_name or not self.site_location_city:
			return

		# First 3 characters of organization name
		customer_code = self.organization_name.strip().lower()[:3]

		# City (remove spaces)
		city = self.site_location_city.strip().lower().replace(" ", "")

		# Creation date (YYYY-MM-DD)
		create_date = (self.creation or frappe.utils.now()).split(" ")[0]

		self.naming_series = f"{customer_code}-{city}-{create_date}"


	def before_validate(self):
		# self.set_sla()
		pass

	def validate(self):
		# self.set_primary_contact()
		# self.set_primary_email_mobile_no()
		# if not self.is_new() and self.has_value_changed("deal_owner") and self.deal_owner:
		# 	self.share_with_agent(self.deal_owner)
		# 	self.assign_agent(self.deal_owner)
		# if self.has_value_changed("status"):
		# 	add_status_change_log(self)
		pass

	def after_insert(self):
		# if self.deal_owner:
		# 	self.assign_agent(self.deal_owner)
		pass

	def before_save(self, method=None):

		if not self.organization_name:
			return

		existing = frappe.db.exists(
			"CRM Organization",
			{
				"organization_name": self.organization_name,
				"project":self.project,
                "branch":self.branch,
				"name": ["!=", self.name]   # current doc ko ignore karega (update case)
			}
		)

		if existing:
			frappe.throw(f"Customer already exists: {self.organization_name}")

	def set_primary_contact(self, contact=None):
		if not self.contacts:
			return

		if not contact and len(self.contacts) == 1:
			self.contacts[0].is_primary = 1
		elif contact:
			for d in self.contacts:
				if d.contact == contact:
					d.is_primary = 1
				else:
					d.is_primary = 0

	def set_primary_email_mobile_no(self):
		if not self.contacts:
			self.email = ""
			self.mobile_no = ""
			self.phone = ""
			return

		if len([contact for contact in self.contacts if contact.is_primary]) > 1:
			frappe.throw(_("Only one {0} can be set as primary.").format(frappe.bold("Contact")))

		primary_contact_exists = False
		for d in self.contacts:
			if d.is_primary == 1:
				primary_contact_exists = True
				self.email = d.email.strip() if d.email else ""
				self.mobile_no = d.mobile_no.strip() if d.mobile_no else ""
				self.phone = d.phone.strip() if d.phone else ""
				break

		if not primary_contact_exists:
			self.email = ""
			self.mobile_no = ""
			self.phone = ""

	def assign_agent(self, agent):
		if not agent:
			return

		assignees = self.get_assigned_users()
		if assignees:
			for assignee in assignees:
				if agent == assignee:
					# the agent is already set as an assignee
					return

		assign({"assign_to": [agent], "doctype": "CRM Organization", "name": self.name})

	def share_with_agent(self, agent):
		if not agent:
			return

		docshares = frappe.get_all(
			"DocShare",
			filters={"share_name": self.name, "share_doctype": self.doctype},
			fields=["name", "user"],
		)

		shared_with = [d.user for d in docshares] + [agent]

		for user in shared_with:
			if user == agent and not frappe.db.exists("DocShare", {"user": agent, "share_name": self.name, "share_doctype": self.doctype}):
				frappe.share.add_docshare(
					self.doctype, self.name, agent, write=1, flags={"ignore_share_permission": True}
				)
			elif user != agent:
				frappe.share.remove(self.doctype, self.name, user)

	def set_sla(self):
		if self.sla:
			return

		sla = get_sla(self)
		if not sla:
			self.first_responded_on = None
			self.first_response_time = None
			return
		self.sla = sla.name

	def apply_sla(self):
		if not self.sla:
			return
		sla = frappe.get_last_doc("CRM Service Level Agreement", {"name": self.sla})
		if sla:
			sla.apply(self)


@frappe.whitelist()
def make_crm_deal(source_name, target_doc=None):
	def set_missing_values(source, target):
		target.date = frappe.utils.nowdate()
		target.valid_till = frappe.utils.add_days(target.date, 30)
		target.status = "Qualification"
		target.series = "QT.YYYY.-"
		target.deal_type = "Warranty AMC Conversion"
	
	doclist = get_mapped_doc(
		"CRM Organization",
		source_name,
		{
			"CRM Organization": {
				"doctype": "CRM Deal",
				"field_map": {
					"Customer_HC": "Customer_HC",
					"organization_name":"customer_name",
					"custom_asm_name":"deal_owner"
				},
				"validation": {
					"docstatus": ["=", 0]
				}
			}
		},
		target_doc,
		set_missing_values
	)
	return doclist

import frappe
from frappe import _
from frappe.utils import now_datetime
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from io import BytesIO

# ----------------------------
# CONFIG
# ----------------------------
TARGET_DOCTYPE = "CRM Organization"
CHILD_TABLE = "customer_visit_form"
CHILD_DOCTYPE = "Customer Visit Form"
DEFAULT_UNIQUE_FIELD = "customer_hc"
AUTO_CREATE_LINKS = False
DEFAULT_PHONE_COUNTRY_CODE = "+91"

LINK_FIELD_TO_DOCTYPE = {
    "region": "Region Master",
    "branch": "Region Branches",
}

# Excel Header -> DocField fieldname aliases (case-insensitive)
ALIASES = {
    # core
    "customer name": "organization_name",
    "organization name": "organization_name",
    "customer hc": "customer_hc",
    "region": "region",
    "branch": "branch",

    # contact
    "email id": "email",
    "email": "email",
    "mail": "email",

    # phone
    "mobile no": "phone",
    "mobile": "phone",
    "phone": "phone",
    "phone no": "phone",
    "contact no": "phone",

    # phone cc (input column)
    "country code": "phone_country_code",
    "phone country code": "phone_country_code",
    "mobile country code": "phone_country_code",
    "is country code": "phone_country_code",
    "phone cc": "phone_country_code",
}

@frappe.whitelist()
def download_excel_template():
    """CLEAN template: header + 1 sample parent + 1 child row ONLY"""
    doctype = TARGET_DOCTYPE
    meta = frappe.get_meta(doctype)
    child_meta = frappe.get_meta(CHILD_DOCTYPE)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = doctype

    # Parent table headers
    parent_headers = [
        "Customer hc",
        "Organization name", 
        "Site Location (City)",     
        "Vertical",
        "Region",
        "Branch",
        "Email id",
        "Country code",
        "Mobile no"
    ]

    # Child table headers - CLEAN LABELS, NO PREFIX ✅
    child_headers = [
        "Customer Visit Date",
        "Description",
        "Accompanied By", 
        "Purpose of Visit",
        "Other Purpose of Visit"
    ]

    # Write parent headers
    for col_idx, header in enumerate(parent_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(header) + 2, 15)

    # Write child headers (starting from next column)
    child_start_col = len(parent_headers) + 2
    for col_idx, header in enumerate(child_headers, start=child_start_col):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(header) + 2, 20)

    # ⭐ Sample row - parent + 1st child
    sample_row1 = [
        "", "ABC Corp Pvt Ltd", "Delhi", "IT", 
        "NORTH-1", "DEL", "contact@abccorp.com", 
        "+91", "9876543210", "", 
        "2026-02-10", "First visit - demo scheduled", 
        "abhijeet.nadkar@lge.com", "AMC Renewal", ""
    ]
    
    for col_idx, value in enumerate(sample_row1, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    ws.freeze_panes = "A5"
    
    filename = f"CRM_Organization_Import_Template_{now_datetime().strftime('%Y%m%d_%H%M%S')}.xlsx"
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    file_doc = frappe.new_doc("File")
    file_doc.file_name = filename
    file_doc.content = output.getvalue()
    file_doc.is_private = 0
    file_doc.folder = "Home"
    file_doc.insert(ignore_permissions=True)
    
    frappe.db.commit()
    return file_doc.file_url

# ----------------------------
# HELPERS
# ----------------------------
def _read_xlsx(file_url: str):
    """Read Excel directly from Frappe File URL - ✅ FIXED"""
    import openpyxl
    from io import BytesIO
    
    try:
        # Get File document by file_url
        file_doc = frappe.get_doc("File", {"file_url": file_url})
        file_content = file_doc.get_content()
        
        wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
        ws = wb.active

        data = []
        for row in ws.iter_rows(values_only=True):
            data.append(list(row))

        # Remove trailing empty rows
        while data and len(data) > 1 and not any(cell is not None and str(cell).strip() for cell in data[-1]):
            data.pop()

        if not data or len(data) < 2:
            frappe.throw(_("Excel must contain header row + at least 1 data row."))

        headers = [str(h).strip() if h is not None else "" for h in data[0]]
        rows = data[1:]
        return headers, rows
        
    except frappe.DoesNotExistError:
        frappe.throw(_("File not found: {0}").format(file_url))
    except Exception as e:
        frappe.throw(_("Cannot read Excel file: {0}").format(str(e)))

def _get_meta_fields(meta):
    return meta.get("fields") or meta.fields or []

def _build_field_resolver(meta):
    label_to_field = {}
    scrub_to_field = {}
    fieldnames = set()

    for df in _get_meta_fields(meta):
        fieldname = getattr(df, "fieldname", None) or (df.get("fieldname") if isinstance(df, dict) else None)
        label = getattr(df, "label", None) or (df.get("label") if isinstance(df, dict) else None)

        if not fieldname:
            continue

        fieldnames.add(fieldname)
        scrub_to_field[frappe.scrub(fieldname)] = fieldname

        if label:
            lbl = str(label).strip()
            label_to_field[lbl.lower()] = fieldname
            scrub_to_field[frappe.scrub(lbl)] = fieldname

    def resolve(header: str):
        h = (header or "").strip()
        if not h:
            return None

        h_lower = h.lower()
        h_scrub = frappe.scrub(h)

        # 0) aliases
        if h_lower in ALIASES:
            return ALIASES[h_lower]

        # 1) label exact
        if h_lower in label_to_field:
            return label_to_field[h_lower]

        # 2) scrubbed label/fieldname
        if h_scrub in scrub_to_field:
            return scrub_to_field[h_scrub]

        # 3) header is fieldname
        if h_scrub in fieldnames:
            return h_scrub

        return None

    return resolve

def _normalize_value(val):
    if val is None:
        return None

    if isinstance(val, str):
        v = val.strip()
        return v if v else None

    if isinstance(val, float) and val.is_integer():
        val = int(val)

    return val

def _normalize_country_code(cc):
    if cc is None:
        return None

    cc = str(cc).strip()
    if not cc:
        return None

    if cc.isdigit():
        cc = f"+{cc}"

    if not cc.startswith("+"):
        cc = f"+{cc}"

    return cc

def _format_phone(cc, number):
    cc = _normalize_country_code(cc) or DEFAULT_PHONE_COUNTRY_CODE

    if number is None:
        return None

    num = str(number).strip()
    if not num:
        return None

    num = num.replace(" ", "").replace("-", "")

    if num.startswith("+"):
        num_digits = "".join([c for c in num if c.isdigit()])
        if len(num_digits) >= 10:
            num = num_digits[-10:]
        else:
            num = num_digits
    else:
        num = "".join([c for c in num if c.isdigit()])
        if num.startswith("91") and len(num) > 10:
            num = num[-10:]

    if len(num) != 10:
        raise frappe.ValidationError(_("Invalid mobile number length: {0}").format(num))

    return f"{cc}-{num}"

def _get_mandatory_fields(meta):
    mandatory = set()
    for df in _get_meta_fields(meta):
        fieldname = getattr(df, "fieldname", None) or (df.get("fieldname") if isinstance(df, dict) else None)
        reqd = getattr(df, "reqd", None) if not isinstance(df, dict) else df.get("reqd")

        if fieldname and reqd:
            mandatory.add(fieldname)
    return mandatory

def _validate_links(doc, meta):
    for df in _get_meta_fields(meta):
        fieldtype = getattr(df, "fieldtype", None) or (df.get("fieldtype") if isinstance(df, dict) else None)
        fieldname = getattr(df, "fieldname", None) or (df.get("fieldname") if isinstance(df, dict) else None)

        if fieldtype != "Link" or not fieldname:
            continue

        value = doc.get(fieldname)
        if not value:
            continue

        options = getattr(df, "options", None) or (df.get("options") if isinstance(df, dict) else None)
        if not options:
            continue

        link_doctype = LINK_FIELD_TO_DOCTYPE.get(fieldname, options)

        if frappe.db.exists(link_doctype, value):
            continue

        if AUTO_CREATE_LINKS:
            new_doc = frappe.new_doc(link_doctype)
            title_field = frappe.get_meta(link_doctype).get_title_field() or "name"
            new_doc.set(title_field, value)
            new_doc.insert(ignore_permissions=True)
        else:
            raise frappe.ValidationError(_("Invalid Link: {0} '{1}' not found").format(link_doctype, value))

def _get_upsert_key(doc_dict: dict, meta):
    valid_cols = set(meta.get_valid_columns())

    if DEFAULT_UNIQUE_FIELD in valid_cols:
        val = doc_dict.get(DEFAULT_UNIQUE_FIELD)
        if val:
            return DEFAULT_UNIQUE_FIELD, val

    org = doc_dict.get("organization_name")
    br = doc_dict.get("branch")
    if org and br:
        return "__fallback__", f"{org}||{br}"

    return None, None

def _find_existing_doc(unique_field, unique_value, doc_dict):
    if unique_field == DEFAULT_UNIQUE_FIELD:
        return frappe.db.get_value(TARGET_DOCTYPE, {unique_field: unique_value}, "name")

    if unique_field == "__fallback__":
        return frappe.db.get_value(
            TARGET_DOCTYPE,
            {"organization_name": doc_dict.get("organization_name"), "branch": doc_dict.get("branch")},
            "name",
        )

    return None

# ----------------------------
# MAIN IMPORT FUNCTION
# ----------------------------
@frappe.whitelist()
def import_excel_crm_organization(file_url: str):
    """FIXED: Multiple child rows + NO cvf_ prefix"""
    try:
        if not file_url:
            frappe.throw(_("File is required."))
        if not file_url.lower().endswith(".xlsx"):
            frappe.throw(_("Only .xlsx files are supported."))

        headers, rows = _read_xlsx(file_url)
        meta = frappe.get_meta(TARGET_DOCTYPE)
        child_meta = frappe.get_meta(CHILD_DOCTYPE)
        valid_fields = set(meta.get_valid_columns())
        child_valid_fields = set(child_meta.get_valid_columns())
        resolve_field = _build_field_resolver(meta)
        resolve_child_field = _build_field_resolver(child_meta)

        inserted = skipped = child_added = 0
        row_errors = []
        last_parent_name = None

        for row_no, row in enumerate(rows, start=2):
            try:
                if not any(cell is not None and str(cell).strip() for cell in row):
                    skipped += 1
                    continue

                parent_dict = {}
                current_child_row = {}

                # ⭐ FIXED: Smart parent/child detection - NO cvf_ prefix needed
                for i, header in enumerate(headers):
                    if i >= len(row): continue
                    val = _normalize_value(row[i])
                    if val is None: continue

                    # PARENT fields first (takes priority)
                    fieldname = resolve_field(header)
                    if fieldname and fieldname in valid_fields:
                        parent_dict[fieldname] = val
                        continue  # Skip child processing

                    # CHILD fields - match against child_meta labels/fieldnames
                    child_field = resolve_child_field(header)
                    if child_field and child_field in child_valid_fields:
                        current_child_row[child_field] = val

                # Phone formatting
                if parent_dict.get("phone"):
                    cc = parent_dict.get("phone_country_code")
                    try:
                        parent_dict["phone"] = _format_phone(cc, parent_dict.get("phone"))
                    except: 
                        pass
                parent_dict.pop("phone_country_code", None)

                has_parent_data = bool(parent_dict.get("organization_name"))
                has_child_data = bool(current_child_row)

                # CASE 1: NEW PARENT ROW
                if has_parent_data:
                    last_parent_name = None
                    if "naming_series" not in parent_dict:
                        parent_dict["naming_series"] = "CRM-ORG-.YYYY.-"

                    doc = frappe.new_doc(TARGET_DOCTYPE)
                    for k, v in parent_dict.items():
                        doc.set(k, v)

                    _validate_links(doc, meta)
                    doc.save()
                    frappe.db.commit()

                    last_parent_name = doc.name
                    inserted += 1

                    # Add child if present
                    if has_child_data:
                        child_doc = doc.append(CHILD_TABLE, current_child_row)
                        doc.save()
                        frappe.db.commit()
                        child_added += 1

                # CASE 2: CHILD-ONLY ROW (continues previous parent)
                elif has_child_data and last_parent_name:
                    doc = frappe.get_doc(TARGET_DOCTYPE, last_parent_name)
                    child_doc = doc.append(CHILD_TABLE, current_child_row)
                    doc.save()
                    frappe.db.commit()
                    child_added += 1

                elif has_child_data and not last_parent_name:
                    row_errors.append(f"Row {row_no}: Child needs parent first")
                    continue
                else:
                    skipped += 1

            except Exception as e:
                row_errors.append(f"Row {row_no}: {str(e)}")
                frappe.db.rollback()

        msg = _("Import Complete!") + "<br>"
        if inserted: msg += f"Parents: {inserted}<br>"
        if child_added: msg += f"Children: {child_added}<br>"
        if skipped: msg += f"Skipped: {skipped}<br>"

        if row_errors:
            msg += "<br><b>Errors:</b><br>" + "<br>".join(row_errors)

        frappe.msgprint(msg)
        return msg

    except Exception as e:
        frappe.throw(_("Import failed: {}").format(str(e)))

# @frappe.whitelist()
# def make_project(source_name, target_doc=None):
# 	def set_missing_values(source, target):
# 		pass
# 		# target.date = frappe.utils.nowdate()
# 		# target.valid_till = frappe.utils.add_days(target.date, 30)
# 		# target.status = "Qualification"
# 		# target.series = "QT.YYYY.-"
	
# 	doclist = get_mapped_doc(
# 		"CRM Organization",
# 		source_name,
# 		{
# 			"CRM Organization": {
# 				"doctype": "Project",
# 				"field_map": {
# 					"organization_name": "customer_name",
# 					"customer_primary_address":"customer_address"
# 				},
# 				"validation": {
# 					"docstatus": ["=", 0]
# 				}
# 			}
# 		},
# 		target_doc,
# 		set_missing_values
# 	)
# 	return doclist



# import frappe
# from frappe.contacts.doctype.address.address import get_address_display

# @frappe.whitelist()
# def get_address_by_org_name(org_name):
#     frappe.logger().info(f"Looking for org: {org_name}")
    
#     org = frappe.get_value("CRM Organization", {"organization_name": org_name}, "name")
#     if not org:
#         return "<div>No organization found with that name</div>"

#     frappe.logger().info(f"Organization ID: {org}")

#     address_links = frappe.get_all(
#         "Dynamic Link",
#         filters={
#             "link_doctype": "CRM Organization",
#             "link_name": org,
#             "parenttype": "Address"
#         },
#         fields=["parent"],
#         limit=1
#     )

#     if not address_links:
#         return "<div>No address linked to this organization</div>"

#     address_name = address_links[0].parent
#     frappe.logger().info(f"Address found: {address_name}")

#     address_html = get_address_display(address_name)

#     # Return HTML with an edit icon (using Font Awesome for example)
#     return f"""
# <div style="border: 1px solid #ccc; border-radius: 6px; padding: 15px; background-color: #f9f9f9; font-size: 14px; position: relative;">
#     <div>{address_html}</div>
#     <a href="javascript:void(0)" 
#        style="position: absolute; top: 10px; right: 10px; cursor: pointer; color: #007bff; font-size: 18px;" 
#        onclick="open_address_and_refresh('{address_name}');" 
#        title="Edit Address">
#        &#9998;
#     </a>
# </div>
# """
